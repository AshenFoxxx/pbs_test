from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.formatters import format_date
from app.metrics import STATUS_LABELS

INK = "1E2430"
HEADER = "0E7C66"
PALE = "F8E6E2"
LINE = "E3DCCB"


def build_report(model: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Итоги"
    _write_summary(summary, model)
    _write_prep(workbook.create_sheet("Заготовка"), model["prep"]["rows"])
    _write_weld(workbook.create_sheet("Сварка"), model["weld"]["rows"])
    _write_docs(workbook.create_sheet("Документация"), model["docs"]["rows"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_summary(sheet, model: dict) -> None:
    sheet["A1"] = "Сводка участков"
    sheet["A1"].font = Font(bold=True, size=16, color=INK)
    sheet["A2"] = f"Период: {format_date(model['date_from'])} — {format_date(model['date_to'])}"
    sheet["A4"] = "Участок"
    sheet["B4"] = "Показатель"
    sheet["C4"] = "Факт"
    sheet["D4"] = "План или потребность"
    sheet["E4"] = "Процент"
    sheet["F4"] = "В итоге"
    sheet["G4"] = "Всего строк"
    _header(sheet, 4, 7)

    rows = [
        ("Заготовка", "Выполнение плана", model["prep"]),
        ("Сварка", "Покрытие потребности", _weld_as_pair(model["weld"])),
        ("Документация", "Выполнение плана", model["docs"]),
    ]
    for offset, (area, title, totals) in enumerate(rows, start=5):
        sheet.cell(offset, 1, area)
        sheet.cell(offset, 2, title)
        sheet.cell(offset, 3, totals["fact"])
        sheet.cell(offset, 4, totals["plan"])
        sheet.cell(offset, 5, totals["ratio"] if totals["ratio"] is not None else "—")
        if totals["ratio"] is not None:
            sheet.cell(offset, 5).number_format = "0.0%"
        sheet.cell(offset, 6, totals["used"])
        sheet.cell(offset, 7, totals["total"])

    sheet["A9"] = "Как посчитано"
    sheet["A9"].font = Font(bold=True, size=13, color=INK)
    for index, rule in enumerate(model["rules"], start=10):
        sheet.cell(index, 1, rule)
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=7)
        sheet.cell(index, 1).alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[index].height = 32

    blocked = model["weld"]["blocked_batches"]
    if blocked:
        row_index = 10 + len(model["rules"]) + 1
        sheet.cell(row_index, 1, "Партии без достоверного покрытия: " + ", ".join(blocked))
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)

    _widths(sheet, [22, 28, 14, 26, 12, 12, 14])
    sheet.row_dimensions[2].height = 18


def _weld_as_pair(weld: dict) -> dict:
    return {
        "fact": weld["fact"],
        "plan": weld["need"],
        "ratio": weld["ratio"],
        "used": weld["used"],
        "total": weld["total"],
    }


def _write_prep(sheet, rows: list[dict]) -> None:
    headers = ["Дата", "Запись", "Партия", "Сотрудник", "План", "Факт", "В итоге", "Почему исключена"]
    _table_header(sheet, headers)
    for index, row in enumerate(rows, start=2):
        values = [
            row["date"],
            row["record_id"],
            row["batch_id"],
            row["employee"],
            _excel_number(row["plan_qty"]),
            _excel_number(row["fact_qty"]),
            "да" if row["reliable"] else "нет",
            ", ".join(row["issues"]),
        ]
        _write_row(sheet, index, values, bad=not row["reliable"])
    _widths(sheet, [14, 12, 12, 22, 12, 12, 12, 42])


def _write_weld(sheet, rows: list[dict]) -> None:
    headers = ["Дата", "Запись", "Партия", "Потребность", "Заготовка", "Факт заготовки", "Статус", "Покрытие", "Пояснение"]
    _table_header(sheet, headers)
    for index, row in enumerate(rows, start=2):
        values = [
            row["date"],
            row["record_id"],
            row["batch_id"],
            _excel_number(row["need_qty"]),
            row["prep_record"] or "—",
            _excel_number(row["prep_fact"]),
            STATUS_LABELS.get(row["status"], row["status"]),
            row["coverage"] if row["coverage"] is not None else "—",
            ", ".join(row["link_issues"] or row["issues"]),
        ]
        _write_row(sheet, index, values, bad=not row["reliable"])
        if row["coverage"] is not None:
            sheet.cell(index, 8).number_format = "0.0%"
    _widths(sheet, [14, 12, 12, 14, 14, 18, 22, 12, 46])


def _write_docs(sheet, rows: list[dict]) -> None:
    headers = ["Дата", "Запись", "Сотрудник", "План", "Факт", "В итоге", "Почему исключена"]
    _table_header(sheet, headers)
    for index, row in enumerate(rows, start=2):
        values = [
            row["date"],
            row["record_id"],
            row["employee"],
            _excel_number(row["plan_docs"]),
            _excel_number(row["fact_docs"]),
            "да" if row["reliable"] else "нет",
            ", ".join(row["issues"]),
        ]
        _write_row(sheet, index, values, bad=not row["reliable"])
    _widths(sheet, [14, 12, 22, 12, 12, 12, 42])


def _table_header(sheet, headers: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        sheet.cell(1, column, header)
    _header(sheet, 1, len(headers))
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    sheet.freeze_panes = "A2"


def _header(sheet, row_index: int, columns: int) -> None:
    fill = PatternFill("solid", fgColor=HEADER)
    font = Font(bold=True, color="FFFFFF")
    for column in range(1, columns + 1):
        cell = sheet.cell(row_index, column)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")


def _write_row(sheet, row_index: int, values: list, bad: bool) -> None:
    fill = PatternFill("solid", fgColor=PALE) if bad else None
    border = Border(
        bottom=Side(style="thin", color=LINE),
    )
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row_index, column, value)
        cell.border = border
        cell.alignment = Alignment(vertical="center")
        if fill is not None:
            cell.fill = fill


def _excel_number(value):
    if value is None:
        return "—"
    return value


def _widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
