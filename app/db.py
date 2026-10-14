import os
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile

ROOT = Path(__file__).resolve().parents[1]

PREP_COLUMNS = ("record_id", "date", "batch_id", "employee", "plan_qty", "fact_qty")
WELD_COLUMNS = ("record_id", "date", "batch_id", "need_qty")
DOCS_COLUMNS = ("record_id", "date", "employee", "plan_docs", "fact_docs")


def db_path() -> Path:
    path = Path(os.environ.get("DB_PATH", ROOT / "data" / "app.db"))
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def seed_dir() -> Path:
    return Path(os.environ.get("SEED_DIR", ROOT / "seed"))


@contextmanager
def connect():
    connection = sqlite3.connect(db_path())
    connection.row_factory = sqlite3.Row
    try:
        init_db(connection)
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def init_db(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS meta (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS prep (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            record_id TEXT NOT NULL,
            date TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            employee TEXT NOT NULL,
            plan_qty REAL,
            fact_qty REAL
        );
        CREATE TABLE IF NOT EXISTS weld (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            record_id TEXT NOT NULL,
            date TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            need_qty REAL
        );
        CREATE TABLE IF NOT EXISTS docs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            record_id TEXT NOT NULL,
            date TEXT NOT NULL,
            employee TEXT NOT NULL,
            plan_docs REAL,
            fact_docs REAL
        );
        """
    )


def database_ready(connection: sqlite3.Connection) -> bool:
    row = connection.execute(
        "SELECT value FROM meta WHERE key = 'seed_imported'"
    ).fetchone()
    return row is not None


def row_counts(connection: sqlite3.Connection) -> dict[str, int]:
    counts = {}
    for table in ("prep", "weld", "docs"):
        counts[table] = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
    return counts


def clear_loaded_data(connection: sqlite3.Connection) -> dict[str, int]:
    counts = row_counts(connection)
    for table in ("prep", "weld", "docs"):
        connection.execute(f"DELETE FROM {table}")
    connection.execute("DELETE FROM meta WHERE key = 'seed_imported'")
    return counts


class UploadRejected(Exception):
    def __init__(self, errors: list[dict]):
        self.errors = errors


FIELD_LABELS = {
    "prep": "Заготовка",
    "weld": "Потребность сварки",
    "docs": "Документация",
}
EXPECTED_FIELDS = {
    "prep": ("record_id", "date", "batch_id", "employee", "plan_qty", "fact_qty"),
    "weld": ("record_id", "date", "batch_id", "need_qty"),
    "docs": ("record_id", "date", "employee", "plan_docs", "fact_docs"),
}


def import_seed(connection: sqlite3.Connection, folder: Path) -> dict[str, int]:
    files = sorted(folder.glob("*.xlsx")) if folder.exists() else []
    if not files:
        raise FileNotFoundError(f"В папке {folder} нет файлов Excel для первой загрузки")
    sources = []
    for file in files:
        kind = _file_kind(file)
        if kind is None:
            raise RuntimeError(f"Файл «{file.name}» не похож на заготовку, сварку или документацию")
        sources.append((kind, file.name, file))
    return import_sources(connection, sources)


def import_sources(connection: sqlite3.Connection, sources: list[tuple[str, str, object]]) -> dict[str, int]:
    if database_ready(connection):
        raise RuntimeError("Данные уже загружены")

    found = {"prep": 0, "weld": 0, "docs": 0}
    seen: dict[str, list[str]] = {kind: [] for kind in found}
    errors = []
    for expected, name, source in sources:
        try:
            headers = _headers_of(source)
        except (InvalidFileException, BadZipFile, ValueError, OSError, KeyError):
            errors.append(_unreadable(expected, name))
        else:
            if _kind(headers) != expected:
                errors.append(_mismatch(expected, name, headers))
            else:
                seen[expected].append(name)
        if hasattr(source, "seek"):
            source.seek(0)
    if errors:
        raise UploadRejected(errors)

    duplicated = [kind for kind, names in seen.items() if len(names) > 1]
    if duplicated:
        raise RuntimeError("Один и тот же участок передан несколькими файлами")
    missing = [kind for kind, names in seen.items() if not names]
    if missing:
        raise RuntimeError("Не хватает данных для загрузки: " + ", ".join(missing))

    for _expected, name, source in sources:
        workbook = load_workbook(source, data_only=False)
        try:
            kinds_in_file = _read_workbook(connection, workbook, found)
        finally:
            workbook.close()
        if not kinds_in_file:
            raise RuntimeError(f"Файл «{name}» не похож на заготовку, сварку или документацию")
    connection.execute(
        "INSERT INTO meta (key, value) VALUES ('seed_imported', ?)",
        (datetime.now().isoformat(timespec="seconds"),),
    )
    return found


def _read_workbook(connection: sqlite3.Connection, workbook, found: dict[str, int]) -> set[str]:
    kinds = set()
    for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [cell for cell in rows[0]]
            kind = _kind(headers)
            if kind is None:
                continue
            index = {name: position for position, name in enumerate(headers)}
            for raw in rows[1:]:
                if raw is None or raw[index["record_id"]] in (None, ""):
                    continue
                if kind == "prep":
                    connection.execute(
                        """
                        INSERT INTO prep (record_id, date, batch_id, employee, plan_qty, fact_qty)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            _text(raw[index["record_id"]]),
                            _date(raw[index["date"]]),
                            _text(raw[index["batch_id"]]),
                            _text(raw[index["employee"]]),
                            _number(raw[index["plan_qty"]]),
                            _number(raw[index["fact_qty"]]),
                        ),
                    )
                elif kind == "weld":
                    connection.execute(
                        """
                        INSERT INTO weld (record_id, date, batch_id, need_qty)
                        VALUES (?, ?, ?, ?)
                        """,
                        (
                            _text(raw[index["record_id"]]),
                            _date(raw[index["date"]]),
                            _text(raw[index["batch_id"]]),
                            _number(raw[index["need_qty"]]),
                        ),
                    )
                else:
                    connection.execute(
                        """
                        INSERT INTO docs (record_id, date, employee, plan_docs, fact_docs)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (
                            _text(raw[index["record_id"]]),
                            _date(raw[index["date"]]),
                            _text(raw[index["employee"]]),
                            _number(raw[index["plan_docs"]]),
                            _number(raw[index["fact_docs"]]),
                        ),
                    )
                found[kind] += 1
                kinds.add(kind)
    return kinds


def fetch(connection: sqlite3.Connection, table: str) -> list[dict]:
    if table not in {"prep", "weld", "docs"}:
        raise ValueError(table)
    cursor = connection.execute(f"SELECT * FROM {table} ORDER BY date, id")
    return [dict(row) for row in cursor]


def update_prep(connection: sqlite3.Connection, row_id: int, data: dict) -> bool:
    cursor = connection.execute(
        """
        UPDATE prep
        SET record_id = ?, date = ?, batch_id = ?, employee = ?, plan_qty = ?, fact_qty = ?
        WHERE id = ?
        """,
        (
            data["record_id"],
            data["date"],
            data["batch_id"],
            data["employee"],
            data["plan_qty"],
            data["fact_qty"],
            row_id,
        ),
    )
    return cursor.rowcount > 0


def insert_prep(connection: sqlite3.Connection, data: dict) -> int:
    cursor = connection.execute(
        """
        INSERT INTO prep (record_id, date, batch_id, employee, plan_qty, fact_qty)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            data["record_id"],
            data["date"],
            data["batch_id"],
            data["employee"],
            data["plan_qty"],
            data["fact_qty"],
        ),
    )
    return int(cursor.lastrowid)


def delete_row(connection: sqlite3.Connection, table: str, row_id: int) -> bool:
    if table not in {"prep", "weld", "docs"}:
        raise ValueError(table)
    cursor = connection.execute(f"DELETE FROM {table} WHERE id = ?", (row_id,))
    return cursor.rowcount > 0


def update_weld(connection: sqlite3.Connection, row_id: int, data: dict) -> bool:
    cursor = connection.execute(
        """
        UPDATE weld
        SET record_id = ?, date = ?, batch_id = ?, need_qty = ?
        WHERE id = ?
        """,
        (data["record_id"], data["date"], data["batch_id"], data["need_qty"], row_id),
    )
    return cursor.rowcount > 0


def insert_weld(connection: sqlite3.Connection, data: dict) -> int:
    cursor = connection.execute(
        """
        INSERT INTO weld (record_id, date, batch_id, need_qty)
        VALUES (?, ?, ?, ?)
        """,
        (data["record_id"], data["date"], data["batch_id"], data["need_qty"]),
    )
    return int(cursor.lastrowid)


def update_docs(connection: sqlite3.Connection, row_id: int, data: dict) -> bool:
    cursor = connection.execute(
        """
        UPDATE docs
        SET record_id = ?, date = ?, employee = ?, plan_docs = ?, fact_docs = ?
        WHERE id = ?
        """,
        (
            data["record_id"],
            data["date"],
            data["employee"],
            data["plan_docs"],
            data["fact_docs"],
            row_id,
        ),
    )
    return cursor.rowcount > 0


def insert_docs(connection: sqlite3.Connection, data: dict) -> int:
    cursor = connection.execute(
        """
        INSERT INTO docs (record_id, date, employee, plan_docs, fact_docs)
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            data["record_id"],
            data["date"],
            data["employee"],
            data["plan_docs"],
            data["fact_docs"],
        ),
    )
    return int(cursor.lastrowid)


def _headers_of(source) -> list[str]:
    workbook = load_workbook(source, data_only=False)
    try:
        recognized = None
        fallback: list[str] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(max_row=1, values_only=True))
            if not rows or rows[0] is None:
                continue
            headers = [str(cell) for cell in rows[0] if cell not in (None, "")]
            if _kind(headers) and recognized is None:
                recognized = headers
            elif len(headers) > len(fallback):
                fallback = headers
        return recognized if recognized is not None else fallback
    finally:
        workbook.close()


def _mismatch(expected: str, filename: str, headers: list[str]) -> dict:
    wanted = EXPECTED_FIELDS[expected]
    missing = [name for name in wanted if name not in headers]
    extra = [name for name in headers if name not in wanted]
    return {
        "field": FIELD_LABELS[expected],
        "filename": filename,
        "expected": ", ".join(wanted),
        "actual": ", ".join(headers) if headers else "нет строки заголовков",
        "missing": ", ".join(missing),
        "extra": ", ".join(extra),
    }


def _unreadable(expected: str, filename: str) -> dict:
    return {
        "field": FIELD_LABELS[expected],
        "filename": filename,
        "expected": ", ".join(EXPECTED_FIELDS[expected]),
        "actual": "файл не открывается как Excel",
        "missing": "",
        "extra": "",
    }


def _file_kind(source) -> str | None:
    workbook = load_workbook(source, data_only=False)
    try:
        found = set()
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(max_row=1, values_only=True))
            if not rows or rows[0] is None:
                continue
            kind = _kind(list(rows[0]))
            if kind:
                found.add(kind)
    finally:
        workbook.close()
    if len(found) != 1:
        return None
    return found.pop()


def _kind(headers: list) -> str | None:
    names = set(headers)
    if "fact_docs" in names:
        return "docs"
    if "need_qty" in names:
        return "weld"
    if "fact_qty" in names and "plan_qty" in names:
        return "prep"
    return None


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _number(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def _date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)[:10]
